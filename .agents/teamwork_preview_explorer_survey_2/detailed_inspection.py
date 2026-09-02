import os
import json
import docx

folder = r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\New folder"
files = [f for f in sorted(os.listdir(folder)) if f.endswith('.docx') and not f.startswith('PP - 05 PRODUCTION')]

report = []

for fname in files:
    path = os.path.join(folder, fname)
    doc = docx.Document(path)
    
    report.append(f"================================================================================")
    report.append(f"TEMPLATE: {fname}")
    report.append(f"================================================================================")
    report.append(f"PARAGRAPHS ({len(doc.paragraphs)} total):")
    for pi, p in enumerate(doc.paragraphs):
        runs_repr = [f"R{ri}: {repr(r.text)}" for ri, r in enumerate(p.runs)]
        report.append(f"  P[{pi}]: {repr(p.text)}")
        if runs_repr:
            report.append(f"      Runs: {', '.join(runs_repr)}")
            
    report.append(f"\nTABLES ({len(doc.tables)} total):")
    for ti, t in enumerate(doc.tables):
        report.append(f"\n  --- Table {ti} ({len(t.rows)} rows x {len(t.columns)} cols) ---")
        for ri, r in enumerate(t.rows):
            report.append(f"    Row {ri}:")
            for ci, c in enumerate(r.cells):
                cell_p_runs = []
                for cpi, cp in enumerate(c.paragraphs):
                    c_runs = [f"r{cri}: {repr(cr.text)}" for cri, cr in enumerate(cp.runs)]
                    cell_p_runs.append(f"p{cpi}[{repr(cp.text)} | {'+'.join(c_runs)}]")
                report.append(f"      Cell {ci}: {' // '.join(cell_p_runs)}")

out_path = r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\.agents\teamwork_preview_explorer_survey_2\detailed_inspection.txt"
with open(out_path, "w", encoding="utf-8") as f:
    f.write("\n".join(report))

print(f"Wrote detailed inspection to {out_path}")
