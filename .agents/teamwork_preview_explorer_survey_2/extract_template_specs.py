import os
import docx

folder = r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\New folder"
files = [f for f in sorted(os.listdir(folder)) if f.endswith('.docx') and not f.startswith('PP - 05 PRODUCTION')]

def analyze_template(fname):
    path = os.path.join(folder, fname)
    doc = docx.Document(path)
    
    print(f"================================================================================")
    print(f"FILE: {fname}")
    print(f"================================================================================")
    
    # 1. Title & Header paragraphs
    print("--- PARAGRAPHS OVERVIEW ---")
    for pi, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t:
            print(f"  P[{pi}]: {t}")
            
    # 2. Tables Overview
    print("\n--- TABLES OVERVIEW ---")
    for ti, table in enumerate(doc.tables):
        print(f"  Table {ti}: {len(table.rows)} rows x {len(table.columns)} cols")
        for ri, r in enumerate(table.rows):
            # Print unique non-empty cells
            cells = []
            seen = set()
            for ci, c in enumerate(r.cells):
                txt = " ".join([cp.text.strip() for cp in c.paragraphs if cp.text.strip()])
                if txt and txt not in seen:
                    seen.add(txt)
                    cells.append(f"[C{ci}]: {txt}")
            if cells:
                print(f"    R{ri}: {' | '.join(cells[:4])}")

for f in files:
    analyze_template(f)
