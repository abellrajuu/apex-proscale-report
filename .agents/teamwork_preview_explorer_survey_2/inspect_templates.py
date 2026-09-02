import os
import json
import docx

folder = r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\New folder"
files = [f for f in sorted(os.listdir(folder)) if f.endswith('.docx')]

out_data = {}

for fname in files:
    path = os.path.join(folder, fname)
    doc = docx.Document(path)
    
    file_info = {
        "filename": fname,
        "paragraphs": [],
        "tables": []
    }
    
    for pi, p in enumerate(doc.paragraphs):
        p_text = p.text
        runs_info = [{"text": r.text, "bold": r.bold, "size": str(r.font.size) if r.font else None} for r in p.runs]
        if p_text.strip():
            file_info["paragraphs"].append({
                "index": pi,
                "text": p_text,
                "runs": runs_info
            })
            
    for ti, t in enumerate(doc.tables):
        t_info = {
            "table_index": ti,
            "rows_count": len(t.rows),
            "cols_count": len(t.columns),
            "rows": []
        }
        for ri, r in enumerate(t.rows):
            r_info = []
            for ci, c in enumerate(r.cells):
                r_info.append({
                    "col": ci,
                    "text": c.text,
                    "paragraphs": [cp.text for cp in c.paragraphs if cp.text.strip()]
                })
            t_info["rows"].append(r_info)
        file_info["tables"].append(t_info)
        
    out_data[fname] = file_info

out_path = r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\.agents\teamwork_preview_explorer_survey_2\template_dump.json"
with open(out_path, "w", encoding="utf-8") as f:
    json.dump(out_data, f, indent=2)

print(f"Successfully dumped {len(out_data)} templates to {out_path}")
