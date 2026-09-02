import os
import docx

folder = r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\New folder"
files = [f for f in sorted(os.listdir(folder)) if f.endswith('.docx') and not f.startswith('PP - 05 PRODUCTION')]

for fname in files:
    path = os.path.join(folder, fname)
    doc = docx.Document(path)
    
    print("="*90)
    print(f"TEMPLATE: {fname}")
    print("="*90)
    
    print(f"PARAGRAPHS ({len(doc.paragraphs)}):")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            runs_str = " | ".join([f"r{ri}:'{r.text}'" for ri, r in enumerate(p.runs)])
            print(f"  P[{i}]: {repr(p.text)}")
            print(f"        Runs: {runs_str}")
            
    print(f"\nTABLES ({len(doc.tables)}):")
    for ti, t in enumerate(doc.tables):
        print(f"  --- Table {ti} ({len(t.rows)} rows x {len(t.columns)} cols) ---")
        for ri, r in enumerate(t.rows):
            # Print each cell
            cell_strs = []
            for ci, c in enumerate(r.cells):
                # cell text with runs
                p_strs = []
                for pi, cp in enumerate(c.paragraphs):
                    if cp.text.strip():
                        r_strs = " + ".join([f"'{cr.text}'" for cr in cp.runs])
                        p_strs.append(f"p{pi}({r_strs})")
                cell_strs.append(f"C{ci}: {' / '.join(p_strs) if p_strs else '[EMPTY]'}")
            print(f"    R{ri}: {' || '.join(cell_strs)}")
    print("\n")
