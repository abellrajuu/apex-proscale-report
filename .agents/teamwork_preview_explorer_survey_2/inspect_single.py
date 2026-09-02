import os
import docx

folder = r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\New folder"

def examine(fname):
    path = os.path.join(folder, fname)
    doc = docx.Document(path)
    print("="*80)
    print(f"FILE: {fname}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"  P[{i}]: {repr(p.text)}")
    print(f"Tables: {len(doc.tables)}")
    for ti, t in enumerate(doc.tables):
        print(f"  --- Table {ti} ({len(t.rows)}x{len(t.columns)}) ---")
        for ri, r in enumerate(t.rows):
            cells = []
            for ci, c in enumerate(r.cells):
                txt = " \\n ".join([p.text.strip() for p in c.paragraphs if p.text.strip()])
                cells.append(f"C{ci}: {txt}")
            print(f"    R{ri}: {' | '.join(cells)}")

if __name__ == "__main__":
    import sys
    fnames = sys.argv[1:] if len(sys.argv) > 1 else ["PP-05_02 Job Traveller Card.docx"]
    for f in fnames:
        examine(f)
