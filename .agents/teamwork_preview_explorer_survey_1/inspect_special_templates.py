import docx
import os

files_to_check = [
    "PP-05_18_A ODD System.docx",
    "PP-05_18_B DD System.docx",
    "PP-05_19 Work Instructions.docx"
]

folder = r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\New folder"

for fname in files_to_check:
    fpath = os.path.join(folder, fname)
    doc = docx.Document(fpath)
    print(f"=== {fname} ===")
    print("Paragraphs count:", len(doc.paragraphs))
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"  P{i}: {p.text}")
    print("Tables count:", len(doc.tables))
    for ti, t in enumerate(doc.tables):
        print(f"  Table {ti}: {len(t.rows)} rows x {len(t.columns)} cols")
        for ri, r in enumerate(t.rows):
            print(f"    R{ri}: {[c.text.replace(chr(10), ' ') for c in r.cells]}")
    # Check headers/footers
    for s_idx, sec in enumerate(doc.sections):
        print(f"  Section {s_idx} Header: {sec.header.paragraphs[0].text if sec.header.paragraphs else 'None'}")
        print(f"  Section {s_idx} Footer: {sec.footer.paragraphs[0].text if sec.footer.paragraphs else 'None'}")
