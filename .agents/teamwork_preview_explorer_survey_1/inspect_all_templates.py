import os
import glob
import docx
import json

folder = r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\New folder"
docx_files = sorted(glob.glob(os.path.join(folder, "*.docx")))

summary = []

for fpath in docx_files:
    fname = os.path.basename(fpath)
    if fname.startswith("PP - 05 PRODUCTION RECORDS"):
        continue
    doc = docx.Document(fpath)
    
    file_info = {
        "filename": fname,
        "paragraphs": [],
        "tables": []
    }
    
    for p_idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt:
            runs_info = [r.text for r in p.runs if r.text]
            file_info["paragraphs"].append({
                "index": p_idx,
                "text": txt,
                "runs": runs_info
            })
            
    for t_idx, t in enumerate(doc.tables):
        table_rows = []
        for r_idx, row in enumerate(t.rows):
            row_cells = []
            for c_idx, cell in enumerate(row.cells):
                cell_text = "\n".join([p.text.strip() for p in cell.paragraphs if p.text.strip()])
                row_cells.append(cell_text)
            table_rows.append(row_cells)
        file_info["tables"].append({
            "index": t_idx,
            "shape": f"{len(t.rows)}x{len(t.columns)}",
            "rows": table_rows
        })
        
    summary.append(file_info)

with open(r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\.agents\teamwork_preview_explorer_survey_1\template_breakdown.json", "w", encoding="utf-8") as f:
    json.dump(summary, f, indent=2, ensure_ascii=False)

print(f"Detailed analysis saved for {len(summary)} files.")
