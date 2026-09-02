import os
import glob
import docx

folder = r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\New folder"
docx_files = sorted(glob.glob(os.path.join(folder, "*.docx")))

print(f"Total docx files found: {len(docx_files)}")

for fpath in docx_files:
    fname = os.path.basename(fpath)
    try:
        doc = docx.Document(fpath)
        p_count = len(doc.paragraphs)
        t_count = len(doc.tables)
        
        # sample text
        texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        first_few = " | ".join(texts[:3]) if texts else "EMPTY"
        
        table_shapes = [f"{len(t.rows)}x{len(t.columns)}" for t in doc.tables]
        
        print(f"\n==========================================")
        print(f"File: {fname}")
        print(f"Paragraphs: {p_count}, Tables: {t_count} (Shapes: {table_shapes})")
        print(f"Sample Headings: {first_few[:120]}")
    except Exception as e:
        print(f"Error reading {fname}: {e}")
