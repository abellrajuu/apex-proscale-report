import os
import docx
import json

new_folder = r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\New folder"

templates = [
    ("PP-05_02 Job Traveller Card.docx", "job_traveller_card"),
    ("PP-05_04 Batching System.docx", "batching_system"),
    ("PP-05_05 Remote Indicator.docx", "remote_indicator"),
    ("PP-05_06 Digital Indicator.docx", "digital_indicator"),
    ("PP-05_07 Weigh Feeder.docx", "weigh_feeder"),
    ("PP-05_08 Crane Scale.docx", "crane_scale"),
    ("PP-05_09 Signal Conditioner.docx", "signal_conditioner"),
    ("PP-05_10 tRIP sAFE.docx", "trip_safe"),
    ("PP-05_11 Vibration Switch.docx", "vibration_switch"),
    ("PP-05_12 ACC mV.docx", "acc_mv"),
    ("PP-05_13 ACC Charge.docx", "acc_charge"),
    ("PP-05_14 Vibration Meter.docx", "vibration_meter"),
    ("PP-05_15 Charge Amplifier.docx", "charge_amplifier"),
    ("PP-05_16 Inprocess.docx", "inprocess"),
    ("PP-05_18_A ODD System.docx", "odd_system"),
    ("PP-05_18_B DD System.docx", "dd_system"),
    ("PP-05_19 Work Instructions.docx", "work_instructions"),
    ("PP-05_20 Performance Index.docx", "performance_index"),
    ("PP-05_21 Performance Delay Analysis.docx", "performance_delay_analysis"),
    ("PP-05_22 Equipment List.docx", "equipment_list"),
    ("PP-05_24 Loss in Weigh Feeder.docx", "loss_in_weigh_feeder"),
]

details = {}

for filename, key in templates:
    filepath = os.path.join(new_folder, filename)
    doc = docx.Document(filepath)
    
    # Extract all paragraphs
    paras = []
    for p_idx, p in enumerate(doc.paragraphs):
        p_text = p.text.strip()
        runs_info = [{"text": r.text, "bold": r.bold, "size": str(r.font.size) if r.font.size else None} for r in p.runs]
        paras.append({
            "index": p_idx,
            "text": p_text,
            "runs": runs_info
        })
        
    # Extract all tables
    tables = []
    for t_idx, t in enumerate(doc.tables):
        rows_data = []
        for r_idx, row in enumerate(t.rows):
            row_cells = []
            for c_idx, cell in enumerate(row.cells):
                # cell paragraphs
                c_paras = [cp.text.strip() for cp in cell.paragraphs if cp.text.strip()]
                row_cells.append({
                    "col_idx": c_idx,
                    "text": cell.text.strip(),
                    "paragraphs": c_paras
                })
            rows_data.append({
                "row_idx": r_idx,
                "cells": row_cells
            })
        tables.append({
            "table_idx": t_idx,
            "rows_count": len(t.rows),
            "cols_count": len(t.columns),
            "rows": rows_data
        })
        
    details[filename] = {
        "key": key,
        "paragraphs_count": len(doc.paragraphs),
        "paragraphs": paras,
        "tables_count": len(doc.tables),
        "tables": tables
    }

with open(r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\.agents\teamwork_preview_explorer_survey_3\detailed_templates_dump.json", "w", encoding="utf-8") as f:
    json.dump(details, f, indent=2)

print(f"Dumped detailed analysis of {len(details)} templates.")
