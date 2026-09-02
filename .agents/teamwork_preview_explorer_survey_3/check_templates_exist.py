import os
import docx

templates_docx_source = r"C:\Users\Admin\Desktop\TEST REPORT PRODUCTION final\New folder"

template_keys = [
    ("PP-05_02 Job Traveller Card.docx", "job_traveller_card"),
    ("PP-05_04 Batching System.docx", "batching_system"),
    ("PP-05_05 Remote Indicator.docx", "remote_indicator"),
    ("PP-05_06 Digital Indicator.docx", "digital_indicator"),
    ("PP-05_07 Weigh Feeder.docx", "weigh_feeder"),
    ("PP-05_08 Crane Scale.docx", "crane_scale"),
    ("PP-05_09 Signal Conditioner.docx", "signal_conditioner"),
    ("PP-05_10 tRIP sAFE.docx", "trip_safe"),
    ("PP-05_11 Vibration Switch.docx", "vibration_switch"),
    ("PP-05_12 ACC mV.docx", "acc_mv"),
    ("PP-05_13 ACC Charge.docx", "acc_charge"),
    ("PP-05_14 Vibration Meter.docx", "vibration_meter"),
    ("PP-05_15 Charge Amplifier.docx", "charge_amplifier"),
    ("PP-05_16 Inprocess.docx", "inprocess"),
    ("PP-05_18_A ODD System.docx", "odd_system"),
    ("PP-05_18_B DD System.docx", "dd_system"),
    ("PP-05_19 Work Instructions.docx", "work_instructions"),
    ("PP-05_20 Performance Index.docx", "performance_index"),
    ("PP-05_21 Performance Delay Analysis.docx", "performance_delay_analysis"),
    ("PP-05_22 Equipment List.docx", "equipment_list"),
    ("PP-05_24 Loss in Weigh Feeder.docx", "loss_in_weigh_feeder"),
]

for filename, key in template_keys:
    path = os.path.join(templates_docx_source, filename)
    doc = docx.Document(path)
    print(f"Key: {key:<28} | File: {filename:<40} | Paras: {len(doc.paragraphs):2d} | Tables: {len(doc.tables):2d}")
