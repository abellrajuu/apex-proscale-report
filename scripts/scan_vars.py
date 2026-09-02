import glob
import re
import os
import docx

system_map = {
    'PP-05_02 Job Traveller Card.docx': 'job_traveller_card.html',
    'PP-05_03 Belt Scale.docx': 'belt_scale.html',
    'PP-05_04 Batching System.docx': 'batching_system.html',
    'PP-05_05 Remote Indicator.docx': 'remote_indicator.html',
    'PP-05_06 Digital Indicator.docx': 'digital_indicator.html',
    'PP-05_07 Weigh Feeder.docx': 'weigh_feeder.html',
    'PP-05_08 Crane Scale.docx': 'crane_scale.html',
    'PP-05_09 Signal Conditioner.docx': 'signal_conditioner.html',
    'PP-05_10 tRIP sAFE.docx': 'trip_safe.html',
    'PP-05_11 Vibration Switch.docx': 'vibration_switch.html',
    'PP-05_12 ACC mV.docx': 'acc_mv.html',
    'PP-05_13 ACC Charge.docx': 'acc_charge.html',
    'PP-05_14 Vibration Meter.docx': 'vibration_meter.html',
    'PP-05_15 Charge Amplifier.docx': 'charge_amplifier.html',
    'PP-05_16 Inprocess.docx': 'inprocess.html',
    'PP-05_17 Miscellaneous Items.docx': 'misc_report.html',
    'PP-05_18_A ODD System.docx': 'odd_system.html',
    'PP-05_18_B DD System.docx': 'dd_system.html',
    'PP-05_19 Work Instructions.docx': 'work_instructions.html',
    'PP-05_20 Performance Index.docx': 'performance_index.html',
    'PP-05_21 Performance Delay Analysis.docx': 'performance_delay_analysis.html',
    'PP-05_22 Equipment List.docx': 'equipment_list.html',
    'PP-05_24 Loss in Weigh Feeder.docx': 'loss_in_weigh_feeder.html'
}

for docx_file, html_file in system_map.items():
    docx_path = os.path.join('templates_docx', docx_file)
    html_path = os.path.join('templates', html_file)
    
    if not os.path.exists(docx_path) or not os.path.exists(html_path):
        continue
        
    try:
        doc = docx.Document(docx_path)
    except Exception as e:
        print(f"Failed to read {docx_file}: {e}")
        continue
        
    # Extract tags
    text_content = ""
    for p in doc.paragraphs: text_content += p.text + " "
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text_content += cell.text + " "
                
    tags = set(re.findall(r'\{\{\s*([a-zA-Z0-9_]+)\s*\}\}', text_content))
    
    # Read HTML
    with open(html_path, 'r', encoding='utf-8') as f:
        html = f.read()
        
    missing = []
    for tag in tags:
        if f'name="{tag}"' not in html and f"name='{tag}'" not in html:
            # Special case for arrays/grids which app.js handles differently
            if not tag.startswith('t_') and not tag.startswith('g_') and not tag.startswith('date'):
                missing.append(tag)
                
    if missing:
        print(f"{html_file} is missing: {', '.join(missing)}")
        # Inject missing inputs
        injection = '<div class="section-box"><div class="section-box-title"><i class="fa-solid fa-plus"></i> Auto-Injected Fields</div><div class="grid-col-2">\n'
        for m in missing:
            injection += f'<div class="input-group"><label>{m.replace("_", " ").title()}</label><input type="text" name="{m}" id="{m}"></div>\n'
        injection += '</div></div>\n'
        
        # Insert before form actions or at end of form
        html = re.sub(r'(<div class="form-actions"|<div style="display: flex; justify-content: flex-end;|<button type="submit")', injection + r'\1', html, count=1)
        
        with open(html_path, 'w', encoding='utf-8') as f:
            f.write(html)
