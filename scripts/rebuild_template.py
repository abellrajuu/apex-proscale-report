import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import copy

# Start from the user's latest clean template!
doc = docx.Document('LATEST TEMPLATE/PP-05_03 Belt Scale.docx')

# 1. Paragraphs (Headers, Footers, Info)
# "REMOVE THE : AFTER RESOLTIIIN IN TEMPALTE AND PUT A - THERE AND ALSO AFTWE KG/M PUT SOME MORE SPACING LIKE ADD 3 SPACES MORE"
doc.paragraphs[3].text = "JOB No.\t: {{job_no}}\t\t\t\t\t\tCUSTOMER: {{customer}}"
doc.paragraphs[4].text = "CAPACITY\t: {{capacity}}\t\t                                                    CONVEYOR NO: {{conveyor_no}}"
doc.paragraphs[5].text = "                                                                                                                                        Version - {{version}}"
doc.paragraphs[9].text = "\tResolution - L: {{res_l}}         S: {{res_s}}      r: {{res_r}}      t: {{res_t}}"
doc.paragraphs[13].text = "      Instrument Used: {{instrument_used}}\t\t\t\t\t\t\t\t                      \t\t   "
doc.paragraphs[16].text = "  Tested by: {{tested_by}}\t                                                       \t\t                         Approved by: {{approved_by}}    "
doc.paragraphs[17].text = "         Date: {{date}}\t\t\t\t\t\t\t\t\t(HOD - PDN)"

for idx in [3, 4, 5, 9, 13, 16, 17]:
    for run in doc.paragraphs[idx].runs:
        run.bold = True
        run.font.size = docx.shared.Pt(11)

# 2. Table 0 (Equipment info)
t0 = doc.tables[0]
t0.cell(1, 0).text = "Model No. {{bs_model}}\n\nSerial No. {{bs_serial}}\n\nREMOTE:\nModel No. {{remote_model}}\nSerial No. {{remote_serial}}"
t0.cell(1, 1).text = "Model No. {{sensor_model}}\n\nSerial No.\n1) {{s1}}\n2) {{s2}}\n3) {{s3}}\n4) {{s4}}"
t0.cell(1, 2).text = "{{tacho_type}} Sensor\n\nModel No. {{tacho_model}}\n\nSerial No. {{tacho_serial}}"
t0.cell(1, 3).text = "Model No. {{jbox_model1}}\nSerial No. {{jbox_serial1}}\n\nModel No. {{jbox_model2}}\nSerial No. {{jbox_serial2}}"

# 3. Table 1 (Short Check)
t1 = doc.tables[1]
t1.cell(0, 0).text = "Check for short between Line & Neutral,\nNeutral & Earth , Line & Earth                   : {{short_check}}"

# 4. Table 2 (Test Performed Matrix)
t2 = doc.tables[2]
for i in range(1, 10):
    t2.cell(i, 2).text = f"{{{{spec_{i}}}}}"
    t2.cell(i, 3).text = f"{{{{act_{i}}}}}"

# 5. Table 3 (Calibration Data Matrix)
t3 = doc.tables[3]
# Add the 7th column for AO2
t3.add_column(Inches(0.9))

# Set fixed widths to prevent page spilling
target_width = Inches(0.9)
for r_idx, row in enumerate(t3.rows):
    for c_idx, cell in enumerate(row.cells):
        cell.width = target_width
        
        # Copy border from column 5 to column 6
        if c_idx == 6:
            c_left = row.cells[5]
            c_right = cell
            try:
                tcPr_left = c_left._tc.get_or_add_tcPr()
                tcPr_right = c_right._tc.get_or_add_tcPr()
                borders_left = tcPr_left.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders')
                if borders_left is not None:
                    existing_borders = tcPr_right.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders')
                    if existing_borders is not None:
                        tcPr_right.remove(existing_borders)
                    tcPr_right.append(copy.deepcopy(borders_left))
            except:
                pass

# Fill Table 3 headers (Row 0)
t3.cell(0, 0).text = "Load on\nLoad cell\nKg / mV"
t3.cell(0, 1).text = "Belt load\nKg/m"
t3.cell(0, 2).text = "Speed in\nm/s"
t3.cell(0, 3).text = "Rate in\ntph"
t3.cell(0, 4).text = "Totalizer in\nTones (6 min)"
t3.cell(0, 5).text = "O/p current mA\nAO1"
t3.cell(0, 6).text = "O/p current mA\nAO2"

# Center align all headers and bold them
for c in range(7):
    p = t3.cell(0, c).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.bold = True

# Fill Table 3 rows 1 to 6
for r in range(1, 7):
    for c in range(5):
        t3.cell(r, c).text = f"{{{{g{r}_{c}}}}}"
        t3.cell(r, c).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    t3.cell(r, 5).text = f"{{{{g{r}_5_ao1}}}}"
    t3.cell(r, 5).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    t3.cell(r, 6).text = f"{{{{g{r}_5_ao2}}}}"
    t3.cell(r, 6).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.save('templates_docx/PP-05_03 Belt Scale.docx')
print("Template rebuilt successfully!")
