import docx

doc = docx.Document('templates_docx/PP-05_02 Job Traveller Card.docx')
found = False
for table in doc.tables:
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            if 'ASSEMBLY' in cell.text:
                print(f"Found ASSEMBLY in table, row {i}, col {j}")
                print("Cell text:", cell.text.replace('\n', ' | '))
                # We will clear this cell and write 'N/A'
                cell.text = 'N/A\n\n'
                # Center and bold
                for p in cell.paragraphs:
                    p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.bold = True
                        r.font.size = docx.shared.Pt(36)
                found = True
if found:
    doc.save('templates_docx/PP-05_02 Job Traveller Card.docx')
    print("Saved modified template!")
