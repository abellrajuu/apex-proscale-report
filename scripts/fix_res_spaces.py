import docx

doc = docx.Document('templates_docx/PP-05_03 Belt Scale.docx')
doc.paragraphs[9].text = "Resolution: L: {{res_l}}      S: {{res_s}}      r: {{res_r}}      t: {{res_t}}"
for r in doc.paragraphs[9].runs:
    r.bold = True
    r.font.size = docx.shared.Pt(11)

doc.save('templates_docx/PP-05_03 Belt Scale.docx')
print("Fixed spaces")
