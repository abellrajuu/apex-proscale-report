import docx
doc = docx.Document('templates_docx/PP-05_06 Digital Indicator.docx')
for p in doc.paragraphs:
    print(p.text.strip())
print("--- TABLES ---")
for t in doc.tables:
    for r in t.rows:
        row_text = []
        for c in r.cells:
            row_text.append(c.text.strip().replace('\n', ' '))
        print(" | ".join(row_text))
