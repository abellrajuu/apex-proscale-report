import glob
import os
import re
import docx

def get_tags_from_docx(doc_path):
    doc = docx.Document(doc_path)
    tags = set()
    for p in doc.paragraphs:
        for match in re.findall(r'\{\{\s*(.*?)\s*\}\}', p.text):
            tags.add(match)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    for match in re.findall(r'\{\{\s*(.*?)\s*\}\}', p.text):
                        tags.add(match)
    return sorted(list(tags))

for doc_path in glob.glob('templates_docx/*.docx'):
    tags = get_tags_from_docx(doc_path)
    print(f"{os.path.basename(doc_path)}: {len(tags)} tags")
