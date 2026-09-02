import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document('templates_docx/PP-05_03 Belt Scale.docx')
t = doc.tables[3]

# Merge cell(0, 5) and cell(0, 6)
t.cell(0, 5).merge(t.cell(0, 6))
t.cell(0, 5).text = ""
p = t.cell(0, 5).paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("O/p current mA\nAO1                    AO2")
run.bold = True

# Fill placeholders for rows 1 to 6
for r in range(1, 7):
    c_left = t.cell(r, 5)
    c_right = t.cell(r, 6)
    
    c_left.text = ""
    p_left = c_left.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_left.add_run(f"{{{{g{r}_5_ao1}}}}")
    
    c_right.text = ""
    p_right = c_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_right.add_run(f"{{{{g{r}_5_ao2}}}}")

doc.save('templates_docx/PP-05_03 Belt Scale.docx')
print("Table 3 AO1/AO2 split completed!")
