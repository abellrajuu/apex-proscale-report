import os
import sys
import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from werkzeug.utils import secure_filename
import threading
import copy

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates_docx")
EXPORTS_DIR = os.path.join(BASE_DIR, "TESTING OUTPUTS")
os.makedirs(EXPORTS_DIR, exist_ok=True)

TEMPLATE_MAP = {
    "job_traveller_card": "PP-05_02 Job Traveller Card.docx",
    "belt_scale": "PP-05_03 Belt Scale.docx",
    "batching_system": "PP-05_04 Batching System.docx",
    "remote_indicator": "PP-05_05 Remote Indicator.docx",
    "digital_indicator": "PP-05_06 Digital Indicator.docx",
    "weigh_feeder": "PP-05_07 Weigh Feeder.docx",
    "crane_scale": "PP-05_08 Crane Scale.docx",
    "signal_conditioner": "PP-05_09 Signal Conditioner.docx",
    "trip_safe": "PP-05_10 tRIP sAFE.docx",
    "vibration_switch": "PP-05_11 Vibration Switch.docx",
    "acc_mv": "PP-05_12 ACC mV.docx",
    "acc_charge": "PP-05_13 ACC Charge.docx",
    "vibration_meter": "PP-05_14 Vibration Meter.docx",
    "charge_amplifier": "PP-05_15 Charge Amplifier.docx",
    "inprocess_register": "PP-05_16 Inprocess.docx",
    "inprocess": "PP-05_16 Inprocess.docx",
    "misc_report": "PP-05_17 Miscellaneous Items.docx",
    "misc": "PP-05_17 Miscellaneous Items.docx",
    "odd_system": "PP-05_18_A ODD System.docx",
    "dd_system": "PP-05_18_B DD System.docx",
    "work_instructions": "PP-05_19 Work Instructions.docx",
    "performance_index": "PP-05_20 Performance Index.docx",
    "performance_delay_analysis": "PP-05_21 Performance Delay Analysis.docx",
    "performance_delay": "PP-05_21 Performance Delay Analysis.docx",
    "equipment_list": "PP-05_22 Equipment List.docx",
    "loss_in_weigh_feeder": "PP-05_24 Loss in Weigh Feeder.docx",
}

def replace_in_paragraph(p, field_map):
    for key, val in field_map.items():
        placeholder = f"{{{{{key}}}}}"
        # Repeat until all occurrences of placeholder are replaced
        while True:
            t_elems = p._p.xpath('.//w:t')
            if not t_elems:
                break
                
            full_text = "".join(t.text for t in t_elems if t.text)
            if placeholder not in full_text:
                break
            
            char_mapping = []
            for t_idx, t in enumerate(t_elems):
                if t.text:
                    for char_idx in range(len(t.text)):
                        char_mapping.append((t_idx, char_idx))
                        
            start_idx = full_text.find(placeholder)
            end_idx = start_idx + len(placeholder) - 1
            
            start_t_idx, start_char_idx = char_mapping[start_idx]
            end_t_idx, end_char_idx = char_mapping[end_idx]
            
            for i in range(start_t_idx + 1, end_t_idx):
                t_elems[i].text = ""
                
            if start_t_idx == end_t_idx:
                r_text = t_elems[start_t_idx].text
                t_elems[start_t_idx].text = r_text[:start_char_idx] + str(val) + r_text[end_char_idx + 1:]
            else:
                r_start = t_elems[start_t_idx].text
                t_elems[start_t_idx].text = r_start[:start_char_idx] + str(val)
                r_end = t_elems[end_t_idx].text
                t_elems[end_t_idx].text = r_end[end_char_idx + 1:]

com_lock = threading.Lock()

def convert_docx_to_pdf(docx_path, pdf_path=None):
    if not pdf_path:
        pdf_path = docx_path.replace(".docx", ".pdf")
    
    with com_lock:
        try:
            import pythoncom
            import win32com.client
            pythoncom.CoInitialize()
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = False
            try:
                word.ActivePrinter = 'Microsoft Print to PDF'
            except:
                pass
            
            doc = word.Documents.Open(os.path.abspath(docx_path))
            doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17) # 17 = wdFormatPDF
            doc.Close(0)
            word.Quit()
        except Exception as e:
            print("Error generating PDF via COM:", e)
    return pdf_path

def generate_docx_record(data, output_path):
    report_type = data.get("report_type", "belt_scale")
    template_name = TEMPLATE_MAP.get(report_type)
    if not template_name:
        raise ValueError(f"Unknown report type: {report_type}")

    template_path = os.path.join(TEMPLATES_DIR, template_name)
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found: {template_path}")

    doc = docx.Document(template_path)

    # Convert all None to empty strings to avoid 'None' printing
    field_map = {k: (v if v is not None else "") for k, v in data.items()}

    # Check for specific hardcoded field defaults that are often missing from UI payload
    if report_type == 'belt_scale':
        if not field_map.get('spec_3'): field_map['spec_3'] = 'Display and Keypad Functionality'
        if not field_map.get('spec_6'): field_map['spec_6'] = 'Number of PF Contacts'
        if not field_map.get('spec_7'): field_map['spec_7'] = 'Communication Output'
        if not field_map.get('spec_8'): field_map['spec_8'] = 'Analog Output'
        if not field_map.get('spec_9'): field_map['spec_9'] = 'Wiring, TB and Component Layout'

    # 1. Paragraph replacement
    for p in doc.paragraphs:
        replace_in_paragraph(p, field_map)

    # 2. Table cell replacement
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_in_paragraph(p, field_map)

    # Output path is passed from app.py
    doc.save(output_path)

    pdf_path = output_path.replace(".docx", ".pdf")
    convert_docx_to_pdf(output_path, pdf_path)

    return pdf_path
